"""
Generate fixture CVs and calibration hire profiles so stages 1-3 can be tested
before Arjun's real applications/ and hires/ folders arrive.

Every person here is invented. The eight hire profiles use the roster printed in
the case brief (names, roles, join dates, ratings) and invent only the narrative
body, which the brief does not supply.

Run:  python scripts/make_fixtures.py
Delete data/applications/* and data/hires/* and drop the real files in to replace.
"""

import json
import os
import pathlib

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

ROOT = pathlib.Path(__file__).resolve().parent.parent
APPS = ROOT / "data" / "applications"
HIRES = ROOT / "data" / "hires"


def slug(name):
    return name.replace(" ", "_")


def write_pdf(path, name, head, body):
    doc = SimpleDocTemplate(str(path), pagesize=A4,
                            topMargin=48, bottomMargin=48,
                            leftMargin=54, rightMargin=54)
    styles = getSampleStyleSheet()
    flow = [Paragraph(f"<b>{name}</b>", styles["Title"])]
    if head:
        flow.append(Paragraph(head, styles["Normal"]))
    flow.append(Spacer(1, 12))
    for line in body:
        if not line.strip():
            flow.append(Spacer(1, 8))
        else:
            flow.append(Paragraph(line.replace("&", "&amp;"), styles["Normal"]))
    doc.build(flow)


def write_docx(path, name, head, body):
    d = Document()
    d.add_heading(name, level=0)
    if head:
        d.add_paragraph(head)
    for line in body:
        d.add_paragraph(line)
    d.save(str(path))


HIRE_PROFILES = [
    {
        "name": "Rohan Desai", "role": "Head of Engineering",
        "joined": "Jul 2022", "rating": "Exceeds Expectations",
        "stood_out": ("Ran a warehouse-management rollout at a 3PL for two years before moving into "
                      "engineering leadership. Described, unprompted, what a dispatch clerk's morning "
                      "looks like."),
        "interview": ("Walked through a migration he aborted three weeks in after the pilot branch's "
                      "numbers did not move. Took the reversal to the client himself."),
        "outcome": "Strongest technical hire to date; sets the bar for the team.",
    },
    {
        "name": "Sunita Krishnamurthy", "role": "Operations Lead",
        "joined": "Jan 2023", "rating": "Exceeds Expectations",
        "stood_out": ("Eight years at a freight forwarder, most of it at the documentation desk and at "
                      "the customs house in person."),
        "interview": ("Had built her own shared tracker for query resolution because the official system "
                      "did not cover it; it was still in use after she left."),
        "outcome": "Runs customer operations end to end. Rarely escalates.",
    },
    {
        "name": "Vikram Nair", "role": "Product Manager",
        "joined": "Jun 2023", "rating": "Meets Expectations",
        "stood_out": ("Four years of PM experience at two well-regarded B2B SaaS companies. Clean roadmap "
                      "artefacts, strong written communication, brand-name employers."),
        "interview": ("Articulate on frameworks and prioritisation. When asked what a freight forwarder "
                      "does between 6am and 9am, could not answer concretely. No example of a decision "
                      "made without a senior sign-off."),
        "outcome": ("Competent and reliable, but needs the problem framed for him. The gap is domain "
                    "contact, not ability."),
    },
    {
        "name": "Aditya Shetty", "role": "Sales Lead",
        "joined": "Aug 2023", "rating": "Exceeds Expectations",
        "stood_out": ("Sold into freight forwarders for five years and had worked a year on the "
                      "operations side before that."),
        "interview": ("Turned down a large deal at his previous company because the customer's process "
                      "would not survive the integration. Made the call alone."),
        "outcome": "Closes the hard accounts. Product trusts his read of the customer.",
    },
    {
        "name": "Preetham Rao", "role": "Backend Engineer",
        "joined": "Feb 2024", "rating": "Below Expectations",
        "stood_out": ("Excellent systems-design interview. Strong pedigree, distributed-systems depth, "
                      "impressive side projects."),
        "interview": ("Every example of hard work was self-contained and technical. No instance of "
                      "changing course based on how the software was actually used."),
        "outcome": ("Builds what is specified, well. Does not notice when the specification is wrong. "
                    "Performance managed."),
    },
    {
        "name": "Meghna Tiwari", "role": "Customer Success Manager",
        "joined": "Aug 2024", "rating": "Exceeds Expectations",
        "stood_out": ("Three years on a transport company's exception desk before moving into customer "
                      "success."),
        "interview": ("Had rewritten her previous employer's escalation runbook on her own initiative "
                      "after tracking where tickets actually stalled; it was adopted."),
        "outcome": "Customers ask for her by name. Feeds the sharpest product input in the company.",
    },
    {
        "name": "Lavanya Iyer", "role": "Product Manager",
        "joined": "Apr 2025", "rating": "Exceeds Expectations",
        "stood_out": ("Two years as an operations analyst at a customs brokerage, then three years in "
                      "product. Less polished on paper than other applicants for the same role."),
        "interview": ("Killed a feature she had personally championed once usage data came in, and told "
                      "the customer who had requested it herself rather than routing it through support."),
        "outcome": ("The benchmark PM hire. Ships things customers adopt without being asked to."),
    },
    {
        "name": "Rahul Bose", "role": "Growth & Marketing Lead",
        "joined": "Jun 2025", "rating": "Meets Expectations",
        "stood_out": ("Strong performance-marketing numbers at a consumer company. No B2B or logistics "
                      "background."),
        "interview": ("Data-driven and decisive on spend, but every decision example sat inside a channel "
                      "dashboard rather than in contact with how customers buy."),
        "outcome": ("Delivers on pipeline targets. Slower than expected to understand the buyer."),
    },
]


def main():
    for d in (APPS / "PM", APPS / "SPM", HIRES):
        d.mkdir(parents=True, exist_ok=True)

    candidates = json.loads((ROOT / "scripts" / "candidates.json").read_text(encoding="utf-8"))
    for c in candidates:
        role = c["role"]
        out = APPS / role / f"{slug(c['name'])}_{role}.{c['fmt']}"
        body = c.get("body", [])
        if c.get("blank"):
            # Stands in for a scanned, image-only CV: parses to almost nothing.
            body = ["[scan]"]
        if c["fmt"] == "pdf":
            write_pdf(out, c["name"], c.get("head", ""), body)
        else:
            write_docx(out, c["name"], c.get("head", ""), body)
        print("CV  ", out.relative_to(ROOT))

    for h in HIRE_PROFILES:
        out = HIRES / f"{slug(h['name'])}.json"
        out.write_text(json.dumps(h, indent=2, ensure_ascii=False), encoding="utf-8")
        print("HIRE", out.relative_to(ROOT))


if __name__ == "__main__":
    main()
